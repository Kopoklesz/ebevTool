"""Egyszeri visszatöltés: személyenkénti munkatörténet a helyi archívumból.

A felhőben tárolt munkatörténet ('ki melyik napon melyik cégnél dolgozott')
csak a legújabb verziótól kezdve létezik. A korábbi hónapok adatai viszont
megvannak a helyi archívumban, a kész statisztika-munkafüzetek 'Név Szerint'
lapján — ez a modul ezekből állítja össze a visszatöltendő adatot.

A modul tiszta Python: nem ír sehová, nem nyúl a Firestore-hoz és a GUI-hoz.
A hívó dönti el, mit kezd az eredménnyel (megmutatja, feltölti stb.).

Az archívum elrendezései a verziótörténet alapján
--------------------------------------------------
  * 86a7d3f (az első, parancssoros változat): még nem volt archívum; a
    '<bemenet>_statisztika.xlsx' a bemeneti fájl mellé került. Ha a
    felhasználó ezeket kézzel gyűjtötte össze, az egy LAPOS mappa
    (cég- és hónapmappa nélkül) — ezt is támogatjuk.
  * 4b96e65 – e20239d / dbcedc6: '<exe mappája>/archívum/<cég>/<ÉÉÉÉ-HH>/'
    (vagy a Beállításokban megadott ARCHIVE_DIR ugyanígy tagolva).
  * fb77301 óta: alapból 'Dokumentumok/ebevTool archívum/<cég>/<ÉÉÉÉ-HH>/';
    ha a régi 'archívum' mappa létezik, továbbra is az a gyökér.
  A cégmappa mindig 'filename_utils.safe_folder_name(cég)', a hónapmappa
  mindig 'generate.serial_to_ym(...)' (ÉÉÉÉ-HH). A kézzel átrendezett
  archívumokhoz a hónapmappa nélküli '<gyökér>/<cég>/<fájl>' alakot is
  elfogadjuk.

A kimeneti munkafüzet formátuma
-------------------------------
A 'Név Szerint' lap blokkszerkezete az első verzió óta változatlan:
    ['név:', név], ['szül.név', ...], ['anyja neve:', ...],
    ['szül.hely, idő:', ...], ['adóazonosító:', adó], ['TAJ-szám:', taj],
    ['lakcím:', ...], üres sor, napi sorok [dátum, 1], ['', darabszám], üres sor.
Ami változott: a személyi adatsorok 3d97b66 előtt (és az e20239d-s ágon)
mindig üresek voltak; az első lap neve a 'e-bev' volt, e20239d óta a bemenet
formátumától függően 'e-bev' vagy 'Bejelentés adatok'. A fájlnév végig
'<bemenet neve>_statisztika.xlsx'.
"""

import os
import re
from collections import Counter
from datetime import date, datetime

from openpyxl import load_workbook

import filename_utils
import generate

NEV_SZERINT = 'Név Szerint'

# A régi és az új archívum-gyökér neve: ha a felhasználó egy szinttel
# feljebb lévő mappát ad meg, ezeket nem tekintjük cégmappának.
ARCHIVE_ROOT_NAMES = ('archívum', 'ebevTool archívum')

# A hónapmappa neve: 'ÉÉÉÉ-HH' (az eltérő elválasztókat is elfogadjuk).
_YM_DIR_RE = re.compile(r'^(\d{4})[-._ ]?(\d{1,2})$')

# A statisztika munkalapjai — ezek közül egyik sem a forrásadat-lap.
_STAT_SHEETS = ('NEV SZERINT', 'DATUM SZERINT', 'KI HANY NAPOT DOLGOZOTT')

# A blokk fejsorainak címkéi (normalizálva) -> mező.
_LABELS = {
    'NEV': 'nev',
    'SZULNEV': 'szul_nev',
    'ANYJANEVE': 'anya_neve',
    'SZULHELYIDO': 'szul_hely_ido',
    'ADOAZONOSITO': 'adoazonosito',
    'ADOAZONOSITOJEL': 'adoazonosito',
    'TAJSZAM': 'taj',
    'TAJ': 'taj',
    'LAKCIM': 'lakcim',
}

# A személyi adatlap mezői (a régebbi kimenetekben üresek lehetnek).
DETAIL_FIELDS = ('szul_nev', 'anya_neve', 'szul_hely_ido', 'lakcim')

# Az Excel-serialként tárolt dátumot csak ebben a tartományban fogadjuk el
# (2000-01-01 .. 2099-12-31), hogy egy darabszámot ne nézzünk dátumnak.
_SERIAL_MIN, _SERIAL_MAX = 36526, 73050


# --- segédfüggvények ---

def _label_key(value):
    """Egy címkecella összehasonlítható alakja: ékezet nélkül, csak betű/szám."""
    if value is None or isinstance(value, (int, float, date)):
        return ''
    return re.sub(r'[^A-Z0-9]', '', filename_utils.normalize(value))


def _is_blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _id_text(value, pad=0):
    """Adóazonosító / TAJ szövegként, akkor is, ha számként tárolták.

    A számként tárolt TAJ elveszti a vezető nullát, ezért azt 'pad'
    hosszúságúra egészítjük ki (a TAJ mindig 9 jegyű).
    """
    if _is_blank(value):
        return ''
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int):
        text = str(value)
        return text.zfill(pad) if pad else text
    return str(value).strip()


def _cell_to_serial(value):
    """Egy dátumcella -> Excel serial, vagy None ha nem dátum.

    Elfogadja a valódi dátumcellát, a szöveges 'ÉÉÉÉ.HH.NN.' alakot és a
    számként maradt Excel-serialt (ha az formázás nélkül került a cellába).
    """
    if isinstance(value, (datetime, date)):
        return generate.date_str_to_serial(value)
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        if float(value).is_integer() and _SERIAL_MIN <= value <= _SERIAL_MAX:
            return int(value)
        return None
    if isinstance(value, str) and value.strip():
        return generate.date_str_to_serial(value)
    return None


def _as_count(value):
    """A darabszám-sor értéke egész számként, vagy None."""
    if isinstance(value, bool) or _is_blank(value):
        return None
    try:
        number = float(str(value).strip().replace(',', '.'))
    except ValueError:
        return None
    return int(number) if number.is_integer() else None


def _find_sheet(wb, wanted):
    """A munkalap a neve alapján; ékezet- és kisbetű-érzéketlen tartalékkal."""
    if wanted in wb.sheetnames:
        return wb[wanted]
    key = _label_key(wanted)
    for name in wb.sheetnames:
        if _label_key(name) == key:
            return wb[name]
    return None


def _open(path):
    return load_workbook(path, read_only=True, data_only=True)


# --- 1. egy munkafüzet beolvasása ---

def _parse_rows(rows, warnings):
    """A 'Név Szerint' lap sorai -> személyek listája (lásd a hívót)."""
    persons = []
    block = None

    def close(block):
        if block is None:
            return
        where = f"{block['row']}. sor"
        nev = block['nev']
        if block['error']:
            warnings.append(f"Hibás blokk kihagyva ({where}, '{nev or '?'}'): {block['error']}")
            return
        if not nev:
            warnings.append(f'Név nélküli blokk kihagyva ({where}).')
            return
        if not block['dates']:
            warnings.append(f"'{nev}' ({where}): nincs egyetlen munkanap sem — kihagyva.")
            return
        unique = sorted(set(block['dates']))
        if len(unique) != len(block['dates']):
            warnings.append(f"'{nev}' ({where}): ismétlődő dátum(ok) — egyszer számolva.")
        if block['count'] is None:
            warnings.append(f"'{nev}' ({where}): hiányzik a darabszám-sor.")
        elif block['count'] != len(unique):
            warnings.append(f"'{nev}' ({where}): a darabszám-sor ({block['count']}) "
                            f'nem egyezik a dátumok számával ({len(unique)}).')
        if not block['adoazonosito'] and not block['taj']:
            warnings.append(f"'{nev}' ({where}): sem adóazonosító, sem TAJ-szám nincs.")
        persons.append({
            'nev': nev,
            'adoazonosito': block['adoazonosito'],
            'taj': block['taj'],
            'dates': [generate.serial_to_iso(s) for s in unique],
            # csak a ténylegesen kitöltött adatlap-mezők
            'details': {f: block[f] for f in DETAIL_FIELDS if block.get(f)},
        })

    for idx, row in enumerate(rows, start=1):
        a = row[0] if len(row) > 0 else None
        b = row[1] if len(row) > 1 else None
        label = _label_key(a)

        if label == 'NEV':
            close(block)
            block = {'row': idx, 'nev': _id_text(b), 'adoazonosito': '', 'taj': '',
                     'dates': [], 'count': None, 'closed': False, 'error': None}
            continue

        if block is None:
            # Az első blokk előtti sorok: üresek vagy ismeretlenek — ez utóbbit jelezzük.
            if not (_is_blank(a) and _is_blank(b)):
                warnings.append(f'{idx}. sor: blokkon kívüli, ismeretlen tartalom — kihagyva.')
            continue

        if _is_blank(a) and _is_blank(b):
            continue  # elválasztó üres sor (akárhány)

        if block['closed']:
            # A darabszám-sor után 'név:' sornak kellene jönnie. Ha más jön,
            # az egy név nélküli (hibás) blokk eleje — külön jelezzük, nem
            # olvasztjuk bele az előzőbe.
            close(block)
            block = {'row': idx, 'nev': '', 'adoazonosito': '', 'taj': '',
                     'dates': [], 'count': None, 'closed': False, 'error': None}

        if block['error']:
            # A blokk már hibás: csak a végét (a darabszám-sort) keressük.
            if _is_blank(a) and _as_count(b) is not None:
                block['closed'] = True
            continue

        if label in _LABELS:
            field = _LABELS[label]
            if field == 'adoazonosito':
                block['adoazonosito'] = _id_text(b)
            elif field == 'taj':
                block['taj'] = _id_text(b, pad=9)
            elif field in DETAIL_FIELDS and not _is_blank(b):
                block[field] = str(b).strip()
            continue

        serial = _cell_to_serial(a)
        if serial is not None:
            block['dates'].append(serial)
            continue

        if _is_blank(a):
            count = _as_count(b)
            if count is not None:
                block['count'] = count
                block['closed'] = True
                continue

        block['error'] = f'{idx}. sor: nem értelmezhető sor ({a!r}, {b!r})'

    close(block)
    return persons


def parse_statistics_workbook(path):
    """Egy statisztika-munkafüzet 'Név Szerint' lapjának beolvasása.

    Visszatérés:
        {'persons': [{'nev': str, 'adoazonosito': str, 'taj': str,
                      'dates': ['ÉÉÉÉ-HH-NN', ...],
                      'details': {mező: érték} — a kitöltött adatlap-mezők
                                 (szul_nev, anya_neve, szul_hely_ido, lakcim)}, ...],
         'warnings': [str, ...]}

    A dátumok rendezettek és egyediek. Pontosan azok a napok, amelyeket az
    adott hónap statisztikája beszámolt — a következő hónapba átnyúló napok
    is, változatlanul. Az adóazonosító és a TAJ mindig szöveg (a számként
    tárolt TAJ 9 jegyre egészül ki).

    Hibás blokkot (név nélküli, értelmezhetetlen sort tartalmazó, dátum
    nélküli) nem dob el kivétellel: kihagyja, és figyelmeztetést ad. Kivételt
    (ValueError) csak akkor dob, ha a fájlban nincs 'Név Szerint' lap — az
    olvashatatlan fájl hibáját az openpyxl kivétele jelzi.
    """
    warnings = []
    wb = _open(path)
    try:
        ws = _find_sheet(wb, NEV_SZERINT)
        if ws is None:
            raise ValueError(f"A fájlban nincs '{NEV_SZERINT}' munkalap: {path}")
        rows = list(ws.iter_rows(max_col=2, values_only=True))
    finally:
        wb.close()
    persons = _parse_rows(rows, warnings)
    persons.sort(key=lambda p: p['nev'].lower())
    return {'persons': persons, 'warnings': warnings}


# --- 2. az archívum bejárása ---

def _ym_from_dirname(name):
    match = _YM_DIR_RE.match(name.strip())
    if not match:
        return None
    year, month = int(match.group(1)), int(match.group(2))
    if not 1 <= month <= 12:
        return None
    return f'{year:04d}-{month:02d}'


def _ym_from_source_sheet(wb):
    """A hónap a beágyazott forrásadat-lapból — pontosan a GUI szabályával.

    A kimenet első lapja a bemeneti fájl teljes tartalma ('e-bev' vagy
    'Bejelentés adatok'). A GUI a hónapot ebből határozta meg
    ('generate.extract_entries': a leggyakoribb bejelentési hónap), így az
    archiválási hónapmappa is ebből lett — ugyanezt számoljuk újra.
    """
    for ws in wb.worksheets:
        if _label_key(ws.title) in _STAT_SHEETS:
            continue
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = rows[0]
        fmt = generate.detect_format(ws.title, header)
        key_col = fmt.col('nev')
        data_rows = [r for r in rows[1:]
                     if len(r) > key_col and r[key_col] and str(r[key_col]).strip()]
        try:
            month_serial = generate.extract_entries(data_rows, fmt)[0]
        except Exception:
            month_serial = None
        if month_serial is not None:
            return generate.serial_to_ym(month_serial)
    return None


def _ym_from_dates(wb):
    """Tartalék: a 'Név Szerint' lap dátumai közül a leggyakoribb hónap.

    A statisztika napjainak túlnyomó része a feldolgozott hónapba esik (csak a
    hónap végén kezdett munkák nyúlnak át a következőbe), ezért a
    leggyakoribb hónap a fájl hónapja. Holtversenynél a korábbi hónap nyer.
    """
    ws = _find_sheet(wb, NEV_SZERINT)
    if ws is None:
        return None
    months = Counter()
    for row in ws.iter_rows(max_col=1, values_only=True):
        serial = _cell_to_serial(row[0] if row else None)
        if serial is not None:
            months[generate.serial_to_ym(serial)] += 1
    if not months:
        return None
    best = max(months.values())
    return min(ym for ym, n in months.items() if n == best)


def _ym_from_filename(path, wb):
    """Utolsó tartalék: hónapnév a fájlnévben, az évet a dátumokból vesszük.

    Év nélkül nem adunk hónapot (a fájl módosítási ideje nem megbízható).
    """
    month = filename_utils.detect_month(path)
    if not month:
        return None
    ws = _find_sheet(wb, NEV_SZERINT)
    years = Counter()
    if ws is not None:
        for row in ws.iter_rows(max_col=1, values_only=True):
            serial = _cell_to_serial(row[0] if row else None)
            if serial is not None:
                years[generate.serial_to_date(serial).year] += 1
    if not years:
        return None
    return f'{years.most_common(1)[0][0]:04d}-{month:02d}'


def _company_and_ym_from_path(rel_dirs):
    """(cégmappa, hónap) a gyökérhez képesti mappalánc alapján.

    '<cég>/<ÉÉÉÉ-HH>' -> (cég, hónap); '<cég>' -> (cég, None);
    gyökér -> (None, None). A hónapmappa a lánc utolsó ilyen eleme, a cég az
    előtte lévő mappa; az archívum-gyökér nevét ('archívum') nem tekintjük
    cégnek.
    """
    ym = None
    company_dirs = list(rel_dirs)
    for i in range(len(rel_dirs) - 1, -1, -1):
        found = _ym_from_dirname(rel_dirs[i])
        if found:
            ym = found
            company_dirs = list(rel_dirs[:i])
            break
    company_dirs = [d for d in company_dirs if d.strip().lower()
                    not in [n.lower() for n in ARCHIVE_ROOT_NAMES]]
    company = company_dirs[-1] if company_dirs else None
    return company, ym


def scan_archive(archive_root, warnings=None):
    """Az archívum összes statisztika-munkafüzetének felderítése.

    Minden elrendezést kezel (lásd a modul leírását): '<gyökér>/<cég>/<ÉÉÉÉ-HH>/',
    '<gyökér>/<cég>/' és lapos mappa. Statisztikának számít minden .xlsx/.xlsm
    fájl, amelyben van 'Név Szerint' lap (így az átnevezett fájlok sem
    vesznek el); az Excel '~$' zárolófájljait kihagyja.

    Visszatérés: [{'path': str, 'company_folder': str|None,
                   'ym': 'ÉÉÉÉ-HH'|None, 'mtime': float}, ...]
    útvonal szerint rendezve.

    A hónap meghatározásának sorrendje:
      1. a hónapmappa neve;
      2. a beágyazott forrásadat-lap leggyakoribb bejelentési hónapja (ez a
         GUI saját szabálya, így ugyanazt adja, amit az archiváláskor);
      3. a 'Név Szerint' dátumai közül a leggyakoribb hónap;
      4. a fájlnévben lévő hónapnév, a dátumok leggyakoribb évével.
    Ha egyik sem sikerül, 'ym' None.

    A 'warnings' listába (ha megadjuk) kerülnek a kihagyott fájlok okai.
    """
    if warnings is None:
        warnings = []
    root = os.path.abspath(archive_root)
    items = []
    if not os.path.isdir(root):
        warnings.append(f'Az archívum mappa nem létezik: {root}')
        return items

    for dirpath, dirnames, filenames in os.walk(root):
        dirnames.sort()
        rel = os.path.relpath(dirpath, root)
        rel_dirs = [] if rel == os.curdir else rel.split(os.sep)
        for filename in sorted(filenames):
            if filename.startswith('~$') or filename.startswith('.~'):
                continue
            if not filename.lower().endswith(('.xlsx', '.xlsm')):
                continue
            path = os.path.join(dirpath, filename)
            company, ym = _company_and_ym_from_path(rel_dirs)
            try:
                mtime = os.path.getmtime(path)
                wb = _open(path)
            except Exception as e:
                warnings.append(f'Nem olvasható, kihagyva: {path} ({e})')
                continue
            try:
                if _find_sheet(wb, NEV_SZERINT) is None:
                    warnings.append(f"Nincs '{NEV_SZERINT}' lapja, kihagyva: {path}")
                    continue
                if ym is None:
                    ym = (_ym_from_source_sheet(wb) or _ym_from_dates(wb)
                          or _ym_from_filename(path, wb))
            except Exception as e:
                warnings.append(f'Nem olvasható, kihagyva: {path} ({e})')
                continue
            finally:
                wb.close()
            items.append({'path': path, 'company_folder': company,
                          'ym': ym, 'mtime': mtime})
    items.sort(key=lambda it: it['path'])
    return items


def candidate_archive_roots(configured=None):
    """Az összes létező archívum-gyökér, ahol a korábbi verziók archiválhattak.

    A Beállításokban megadott mappa ('configured'), a régi '<exe mappája>/
    archívum' és az új 'Dokumentumok/ebevTool archívum' — ismétlés nélkül,
    csak a ténylegesen létezők. Ugyanaz a fájl két gyökérben is lehet: a
    'build_import' ezt a cég+hónap szerinti legújabb-kiválasztással kezeli.
    """
    paths = [configured,
             os.path.join(filename_utils.app_dir(), 'archívum'),
             os.path.join(filename_utils.documents_dir(), 'ebevTool archívum')]
    roots, seen = [], set()
    for p in paths:
        if not p or not os.path.isdir(p):
            continue
        key = os.path.normcase(os.path.abspath(p))
        if key not in seen:
            seen.add(key)
            roots.append(p)
    return roots


# --- 3. cégek megfeleltetése ---

def _loose_key(text):
    return re.sub(r'[^A-Z0-9]+', ' ', filename_utils.normalize(text)).strip()


def match_company_folder(folder_name, known_companies):
    """A cégmappa nevéhez tartozó ismert cégnév, vagy None.

    Elsőként a GUI szabálya: 'safe_folder_name(cég) == mappanév'. Ha ez nem
    talál, kis/nagybetű-, ékezet- és írásjel-érzéketlenül hasonlít. Ha a laza
    összevetés több céget is ad, None (nem találgatunk).
    """
    if not folder_name:
        return None
    companies = [c for c in (known_companies or []) if c]
    for company in companies:
        if filename_utils.safe_folder_name(company) == folder_name:
            return company
    key = _loose_key(folder_name)
    if not key:
        return None
    hits = {c for c in companies
            if _loose_key(filename_utils.safe_folder_name(c)) == key or _loose_key(c) == key}
    return hits.pop() if len(hits) == 1 else None


# --- 4. a visszatöltés összeállítása ---

def build_import(items, company_resolver):
    """A 'scan_archive' eredményéből cég+hónap szerinti visszatöltési adat.

    'company_resolver(company_folder) -> cégnév vagy None' — pl.
    lambda f: match_company_folder(f, ismert_cégek).

    Egy cég+hónaphoz több fájl is tartozhat (újrafeltöltött, javított fájl más
    néven). Ilyenkor KIZÁRÓLAG a legújabbat (módosítási idő szerint) használjuk:
    a javított fájl felülírja a korábbit, így a belőle törölt személy nem
    számít bele. A régebbieket meg sem nyitjuk, csak felsoroljuk. Ha a
    legújabb nem olvasható, a hónap kimarad — nem esünk vissza egy régebbi,
    elavult fájlra.

    Visszatérés:
        {'imports': {(cég, 'ÉÉÉÉ-HH'): {'source': útvonal,
                                         'mtime': float,
                                         'superseded': [útvonal, ...],
                                         'persons': [...]},  # mint a parse-nál
                     ...},
         'unresolved': [{...elem a scan_archive-ból..., 'reason': str}, ...],
         'warnings': [str, ...]}
    """
    warnings = []
    unresolved = []
    groups = {}

    for item in items:
        if not item.get('ym'):
            unresolved.append(dict(item, reason='a hónap nem határozható meg'))
            continue
        try:
            company = company_resolver(item.get('company_folder'))
        except Exception as e:
            company = None
            warnings.append(f"A cég feloldása hibát adott ({item.get('company_folder')!r}): {e}")
        if not company:
            folder = item.get('company_folder')
            reason = (f"ismeretlen cégmappa: '{folder}'" if folder
                      else 'nincs cégmappa (lapos archívum)')
            unresolved.append(dict(item, reason=reason))
            continue
        groups.setdefault((company, item['ym']), []).append(item)

    imports = {}
    for key in sorted(groups):
        company, ym = key
        # Legújabb elöl; azonos időnél az útvonal dönt, hogy determinisztikus legyen.
        candidates = sorted(groups[key], key=lambda it: (it['mtime'], it['path']), reverse=True)
        newest, older = candidates[0], candidates[1:]
        for it in older:
            warnings.append(f'{company} {ym}: régebbi fájl figyelmen kívül hagyva '
                            f"(van újabb: {os.path.basename(newest['path'])}): {it['path']}")
        try:
            parsed = parse_statistics_workbook(newest['path'])
        except Exception as e:
            unresolved.append(dict(newest, reason=f'nem olvasható: {e}'))
            warnings.append(f"{company} {ym}: a legújabb fájl nem olvasható, a hónap kimarad: "
                            f"{newest['path']} ({e})")
            continue
        for w in parsed['warnings']:
            warnings.append(f"{company} {ym} ({os.path.basename(newest['path'])}): {w}")
        imports[key] = {
            'source': newest['path'],
            'mtime': newest['mtime'],
            'superseded': [it['path'] for it in older],
            'persons': parsed['persons'],
        }

    return {'imports': imports, 'unresolved': unresolved, 'warnings': warnings}


# --- 5. személyi adatlapok kinyerése ---

def collect_person_details(imports):
    """A cégenkénti személyek az archívumból.

    'imports': a build_import eredményének 'imports' része. Mindenki benne
    van, aki a cég valamelyik (hónaponként legújabb) kimenetében szerepel —
    akkor is, ha az adatlapja üres volt, hiszen ennél a cégnél dolgozott.
    Személyenként a LEGÚJABB hónap kitöltött adatai számítanak (mezőnként:
    ha az újabb hónapban egy mező üres, egy korábbi hónap értéke pótolja).

    Visszatérés: {cég: [{'nev', 'adoazonosito', 'taj', 'details': {...}}]} —
    adóazonosító és TAJ nélküli személy nem kerül bele (nem párosítható).
    """
    out = {}
    for (company, ym) in sorted(imports, key=lambda k: k[1], reverse=True):
        people = out.setdefault(company, {})
        for p in imports[(company, ym)]['persons']:
            if generate.ado_key(p['adoazonosito']):
                key = generate.ado_key(p['adoazonosito'])
            elif generate.taj_key(p['taj']).isdigit():
                key = 'taj:' + generate.taj_key(p['taj'])
            else:
                continue
            cur = people.setdefault(key, {'nev': p['nev'], 'adoazonosito': p['adoazonosito'],
                                          'taj': p['taj'], 'details': {}})
            for field, value in (p.get('details') or {}).items():
                cur['details'].setdefault(field, value)
            cur['adoazonosito'] = cur['adoazonosito'] or p['adoazonosito']
            cur['taj'] = cur['taj'] or p['taj']
    return {c: list(v.values()) for c, v in out.items() if v}
