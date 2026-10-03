import hashlib
import re
from datetime import date, datetime, timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

from normalize import normalize_taj

EPOCH = datetime(1899, 12, 30)

# --- támogatott bemeneti formátumok ---
#
# A bemeneti fájl kétféle lehet:
#   * 'legacy'  – a régi 'e-bev' munkalap
#   * 'nav2026' – az új NAV-export ('Bejelentés adatok' munkalap,
#                 pl. Egyszerusitett_<adószám>_<riportazonosító>.xlsx)
#
# Az oszlopfelismerés továbbra is pozíció alapú, de a formátumot a munkalap
# neve és a fejléc tartalma alapján automatikusan felismerjük.


class InputFormat:
    """Egy bemeneti formátum pozíció alapú oszlopkiosztása.

    A 'torles' és 'hiba' oszlopok kétféleképp működhetnek:
      * marker módban a cellában szereplő szó jelzi az állapotot (régi fájl),
      * expected módban a cella normál esetben egy fix értéket tartalmaz, és
        minden ettől eltérő érték jelenti a törlést / hibát (új NAV-export,
        ahol nem ismert előre a törölt vagy hibás sorok pontos szövege).
    """

    def __init__(self, key, label, sheet_names, header_keywords, columns,
                 torles=None, hiba=None, output_sheet=None):
        self.key = key
        self.label = label
        self.sheet_names = sheet_names
        self.output_sheet = output_sheet or sheet_names[0]
        self.header_keywords = header_keywords
        self.columns = columns
        self.torles = torles
        self.hiba = hiba

        used = list(columns.values())
        for rule in (torles, hiba):
            if rule:
                used.append(rule['col'])
        self.used_columns = tuple(sorted(set(used)))
        self.min_columns = max(self.used_columns) + 1

    def col(self, name):
        return self.columns[name]


LEGACY_FORMAT = InputFormat(
    key='legacy',
    label="régi 'e-bev' formátum",
    sheet_names=('e-bev',),
    header_keywords=(),
    columns={'nev': 0, 'adoazonosito': 1, 'bejelentes': 2,
             'kezdes': 5, 'munkanapok': 7, 'taj': 8},
    torles={'col': 3, 'mode': 'marker', 'markers': ('törlés',)},
    hiba={'col': 9, 'mode': 'marker', 'markers': ('hibás', 'hiba')},
)

NAV2026_FORMAT = InputFormat(
    key='nav2026',
    label="új NAV-export ('Bejelentés adatok')",
    sheet_names=('bejelentés adatok',),
    header_keywords=('munkavállaló neve', 'adóazonosító jele', 'taj száma',
                     'munkanapok száma', 'bejelentés jellege'),
    # A oszlop = sorszám (nem használjuk), az adatok B-től kezdődnek.
    columns={'nev': 1, 'adoazonosito': 2, 'taj': 3,
             'kezdes': 5, 'munkanapok': 7, 'bejelentes': 13},
    # I oszlop: 'Bejelentés jellege' – normál esetben 'Új'; minden más
    # (törlés, visszavonás, módosítás) törölt rekordnak számít.
    torles={'col': 8, 'mode': 'expected', 'expected': ('új',)},
    # K oszlop: 'Adatlap feldolgozottsági státusza' – normál esetben
    # 'FELDOLGOZOTT'; minden más érték hibás rekordot jelöl.
    hiba={'col': 10, 'mode': 'expected', 'expected': ('feldolgozott',)},
    output_sheet='Bejelentés adatok',
)

FORMATS = (NAV2026_FORMAT, LEGACY_FORMAT)
DEFAULT_FORMAT = LEGACY_FORMAT


_DATE_RE = re.compile(r'^\s*(\d{4})\s*[.\-/]\s*(\d{1,2})\s*[.\-/]\s*(\d{1,2})')


def date_str_to_serial(value):
    """Egy dátum cellaérték -> Excel serial, vagy None ha nem értelmezhető.

    Elfogadja a szöveges 'ÉÉÉÉ.HH.NN.' és 'ÉÉÉÉ-HH-NN' alakot (utána időponttal
    is), valamint a valódi Excel-dátumcellát (datetime/date), mert az
    openpyxl a dátumként formázott cellát már datetime-ként adja vissza.
    """
    if value is None:
        return None
    if isinstance(value, date):  # a datetime is ide tartozik
        return (datetime(value.year, value.month, value.day) - EPOCH).days
    match = _DATE_RE.match(str(value))
    if not match:
        return None
    try:
        y, m, d = (int(g) for g in match.groups())
        return (datetime(y, m, d) - EPOCH).days
    except ValueError:
        return None


def serial_to_date(serial):
    return EPOCH + timedelta(days=serial)


def serial_to_month_serial(serial):
    dt = serial_to_date(serial)
    return (datetime(dt.year, dt.month, 1) - EPOCH).days


def serial_to_ym(serial):
    dt = serial_to_date(serial)
    return f"{dt.year:04d}-{dt.month:02d}"


def serial_to_iso(serial):
    return serial_to_date(serial).strftime('%Y-%m-%d')


def iso_to_serial(iso):
    return (datetime.fromisoformat(iso) - EPOCH).days


def taj_key(taj):
    """A TAJ-szám összehasonlítható alakja: csak a számjegyek.

    Így a '123 456 789' és a '123456789' ugyanannak a személynek számít.
    """
    digits = ''.join(c for c in str(taj or '') if c.isdigit())
    return digits or str(taj or '').strip()


# Az adóazonosítóra ugyanaz a normalizálás érvényes.
ado_key = taj_key


def find_person(persons, adoazonosito, taj):
    """(kulcs, személy) a megadott dolgozóhoz, vagy (None, None).

    Az elsődleges azonosító az adóazonosító: azt ritkábban írják el, mint a
    TAJ-számot. A TAJ csak tartalék — az adóazonosító nélkül rögzített (régi)
    adatlapokhoz, illetve ha a dolgozónak nincs adóazonosítója. Ha egy
    adatlapon VAN adóazonosító, de más, mint a keresett, az másik személy,
    akkor is, ha a TAJ egyezik.
    """
    a_key, t_key = ado_key(adoazonosito), taj_key(taj)
    if a_key:
        for key, p in persons.items():
            if ado_key(p.get('adoazonosito')) == a_key:
                return key, p
    if t_key:
        for key, p in persons.items():
            p_ado = ado_key(p.get('adoazonosito'))
            if taj_key(p.get('taj')) == t_key and (not p_ado or not a_key):
                return key, p
    return None, None


def entry_key(entry):
    return f"{entry['adoazonosito']}|{serial_to_iso(entry['start_serial'])}|{entry['munkanapok']}"


def autofit(ws):
    DATE_FMT = 'YYYY.MM.DD.'
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is None:
                continue
            if isinstance(cell.value, datetime):
                cell.number_format = DATE_FMT
                length = 12
            else:
                length = len(str(cell.value))
            if length > max_len:
                max_len = length
        ws.column_dimensions[col_letter].width = max(max_len + 2, 8)


def cell_text(row, idx):
    if idx is None or len(row) <= idx or row[idx] is None:
        return ''
    return str(row[idx]).strip()


def _flag_value(row, rule):
    """A törlés/hiba szabály által jelzett érték, vagy '' ha a sor rendben van."""
    if not rule:
        return ''
    value = cell_text(row, rule['col'])
    if rule['mode'] == 'marker':
        return value if value.lower() in rule['markers'] else ''
    # expected mód: üres cellát nem tekintünk eltérésnek (hiányzó adat, nem hiba)
    if not value:
        return ''
    return '' if value.lower() in rule['expected'] else value


def torles_value(row, fmt=DEFAULT_FORMAT):
    return _flag_value(row, fmt.torles)


def hiba_value(row, fmt=DEFAULT_FORMAT):
    return _flag_value(row, fmt.hiba)


def is_torles(row, fmt=DEFAULT_FORMAT):
    return bool(torles_value(row, fmt))


def is_hibas(row, fmt=DEFAULT_FORMAT):
    return bool(hiba_value(row, fmt))


# --- visszavonások ---
#
# Nem tudjuk biztosan, hogyan jelenik meg a NAV-exportban egy visszavont
# bejelentés: (a) csak a visszavonó sor marad, (b) az eredeti 'Új' sor is
# megmarad mellette, vagy (c) a visszavonás egy későbbi havi exportban jön.
# A szabály mindhárom esetben helyes: egy (adóazonosító, kezdő nap) kulcsú
# bejelentés érvénytelen, ha UTÁNA ugyanerre a kulcsra visszavonás jött. (a)
# esetben nincs mit párosítani, így semmi nem változik; a visszavonás utáni
# újbóli 'Új' bejelentés pedig érvényes — a visszavonás előtti viszont nem,
# akkor sem, ha más munkanapszámmal jelentették be.
#
# A régi 'e-bev' formátumban a 'törlés' jelölés magát a sort jelöli töröltnek
# (nem egy másik bejelentést von vissza), ezért ott nincs visszavonási esemény.

WITHDRAW_WORDS = ('visszavon', 'törl', 'torl')

_TIME_RE = re.compile(r'(\d{1,2}):(\d{2})(?::(\d{2}))?')


def _event_time(value):
    """A bejelentés időpontja (datetime) a rendezéshez, vagy None."""
    if isinstance(value, datetime):
        return value
    serial = date_str_to_serial(value)
    if serial is None:
        return None
    dt = serial_to_date(serial)
    text = str(value)
    date_match = _DATE_RE.match(text)
    match = _TIME_RE.search(text, date_match.end()) if date_match else None
    if match:
        h, m, s = (int(g or 0) for g in match.groups())
        try:
            dt = dt.replace(hour=h, minute=m, second=s)
        except ValueError:
            pass
    return dt


def row_event(row, fmt=DEFAULT_FORMAT):
    """A sor eseménytípusa: 'new' (érvényes bejelentés), 'withdraw'
    (visszavonás/törlés), 'other' (pl. módosítás, vagy a régi formátum
    törölt sora) vagy None (hibás sor)."""
    if is_hibas(row, fmt):
        return None
    flag = torles_value(row, fmt)
    if not flag:
        return 'new'
    if fmt.torles and fmt.torles['mode'] == 'marker':
        return 'other'
    lowered = flag.lower()
    return 'withdraw' if any(w in lowered for w in WITHDRAW_WORDS) else 'other'


def withdrawal_key(adoazonosito, start_serial, taj=''):
    """A bejelentés kulcsa: adóazonosító (híján TAJ) + kezdő nap, vagy None."""
    ident = ado_key(adoazonosito) or (('taj:' + taj_key(taj)) if taj_key(taj) else '')
    return f'{ident}|{start_serial}' if ident and start_serial is not None else None


def _parse_munkanapok(value):
    try:
        n = int(float(str(value).strip().replace(',', '.')))
        return n if n >= 1 else None
    except (TypeError, ValueError):
        return None


def _row_sort_key(row, fmt, order, index):
    """Rendezőkulcs: bejelentés ideje, majd a fájl (hónap) és a sor sorrendje."""
    when = _event_time(_cell_value(row, fmt.col('bejelentes')))
    return (when or datetime.min, order, index)


def row_events(data_rows, fmt=DEFAULT_FORMAT, order=0):
    """A sorok eseményei: (rendezőkulcs, kulcs, típus, név, kezdő serial, munkanapok)."""
    events = []
    for i, r in enumerate(data_rows):
        kind = row_event(r, fmt)
        if kind not in ('new', 'withdraw'):
            continue
        start = date_str_to_serial(_cell_value(r, fmt.col('kezdes')))
        key = withdrawal_key(cell_text(r, fmt.col('adoazonosito')), start,
                             cell_text(r, fmt.col('taj')))
        if key is None:
            continue
        events.append((_row_sort_key(r, fmt, order, i), key, kind,
                       cell_text(r, fmt.col('nev')), start,
                       _parse_munkanapok(_cell_value(r, fmt.col('munkanapok')))))
    return events


def withdrawal_cutoffs(events):
    """{kulcs: (a legutolsó visszavonás rendezőkulcsa, név, kezdő serial)} —
    a kulcs minden ennél korábbi bejelentése érvénytelen."""
    cutoffs = {}
    for sort_key, key, kind, nev, start, _ in events:
        if kind == 'withdraw' and (key not in cutoffs or sort_key > cutoffs[key][0]):
            cutoffs[key] = (sort_key, nev, start)
    return cutoffs


def valid_registrations(events):
    """Hónapokon átívelő kiértékelés (munkanapló): {kulcs: munkanapszámok}
    azokra a kulcsokra, amikre volt visszavonás — az utolsó visszavonás után
    újra bejelentett munkanapszámok (üres halmaz: végleg visszavonva)."""
    cutoffs = withdrawal_cutoffs(events)
    valid = {key: set() for key in cutoffs}
    for sort_key, key, kind, nev, start, munkanapok in events:
        if key in cutoffs and kind == 'new' and sort_key > cutoffs[key][0]:
            valid[key].add(munkanapok or 1)
    return valid


def entry_withdrawal_key(entry):
    return withdrawal_key(entry['adoazonosito'], entry['start_serial'], entry.get('taj', ''))


def is_withdrawn(entry, withdrawn):
    """Igaz, ha a bejegyzés kulcsára van visszavonás a megadott kulcsok közt
    (egy korábbi fájlból áthozott bejegyzéshez: azt bármely itteni visszavonás
    érvényteleníti, hisz a visszavonás később történt)."""
    return entry_withdrawal_key(entry) in withdrawn


def detect_format(sheet_title, header):
    """A munkalap neve és a fejléc alapján felismert bemeneti formátum."""
    title = (sheet_title or '').strip().lower()
    for fmt in FORMATS:
        if title in fmt.sheet_names:
            return fmt

    header_text = ' | '.join(str(h).strip().lower() for h in (header or []) if h)
    for fmt in FORMATS:
        if fmt.header_keywords and all(kw in header_text for kw in fmt.header_keywords):
            return fmt
    return DEFAULT_FORMAT


def _pick_sheet(wb_in):
    by_title = {ws.title.strip().lower(): ws for ws in wb_in.worksheets}
    for fmt in FORMATS:
        for name in fmt.sheet_names:
            if name in by_title:
                return by_title[name]
    return wb_in.active


def read_input(input_path):
    wb_in = load_workbook(input_path)
    ws_in = _pick_sheet(wb_in)
    all_rows = list(ws_in.values)
    if not all_rows:
        return None, [], DEFAULT_FORMAT
    header = all_rows[0]
    fmt = detect_format(ws_in.title, header)
    key_col = fmt.col('nev')
    data_rows = [r for r in all_rows[1:]
                 if len(r) > key_col and r[key_col] and str(r[key_col]).strip()]
    return header, data_rows, fmt


def check_header(header, fmt=DEFAULT_FORMAT):
    if header is None:
        return ['A fejléc sor hiányzik (üres fájl).']
    problems = []
    if len(header) < fmt.min_columns:
        problems.append(
            f'A fejléc csak {len(header)} oszlopot tartalmaz, legalább {fmt.min_columns} szükséges.')
    for idx in fmt.used_columns:
        if idx < len(header):
            val = header[idx]
            if val is None or not str(val).strip():
                problems.append(f'A fejléc {idx + 1}. oszlopa üres.')
    return problems


def _cell_value(row, idx):
    return row[idx] if idx is not None and len(row) > idx else None


def extract_entries(data_rows, fmt=DEFAULT_FORMAT, problems=None, withdrawn=None):
    """A feldolgozandó bejegyzések kinyerése.

    A 'problems' listába (ha megadjuk) kerül minden sor, amit nem lehetett
    rendesen értelmezni, (név, leírás) párként — ezeket a hívónak kell
    jeleznie, különben a statisztika csendben hiányos lenne.

    A 'withdrawn' szótárba (ha megadjuk) kerülnek a fájlban szereplő
    visszavonások kulcsai {kulcs: (név, kezdő serial)} — a hívó ezekkel a
    korábbi hónapokból átvitt rekordokat szűri. Az ugyanebben a fájlban a
    visszavonás ELŐTT tett 'Új' bejelentések itt kimaradnak.
    """
    if problems is None:
        problems = []
    active = [(i, r) for i, r in enumerate(data_rows)
              if not is_torles(r, fmt) and not is_hibas(r, fmt)]
    active_rows = [r for _, r in active]
    cutoffs = withdrawal_cutoffs(row_events(data_rows, fmt))
    if withdrawn is not None:
        withdrawn.update({k: (nev, start) for k, (_, nev, start) in cutoffs.items()})

    c_bejelentes = fmt.col('bejelentes')
    c_kezdes = fmt.col('kezdes')
    c_munkanapok = fmt.col('munkanapok')

    bejelentes_months = {}
    for r in active_rows:
        s = date_str_to_serial(_cell_value(r, c_bejelentes))
        if s:
            ms = serial_to_month_serial(s)
            bejelentes_months[ms] = bejelentes_months.get(ms, 0) + 1
    current_month_serial = max(bejelentes_months, key=bejelentes_months.get) if bejelentes_months else None

    current_entries = []
    future_entries = []
    for i, r in active:
        nev = cell_text(r, fmt.col('nev'))
        start_serial = date_str_to_serial(_cell_value(r, c_kezdes))
        if start_serial is None:
            problems.append((nev, f'értelmezhetetlen kezdő dátum '
                                  f'("{cell_text(r, c_kezdes)}") — a sor KIMARADT'))
            continue
        try:
            munkanapok = int(float(str(_cell_value(r, c_munkanapok)).strip().replace(',', '.')))
            if munkanapok < 1:
                raise ValueError
        except (TypeError, ValueError):
            problems.append((nev, f'értelmezhetetlen munkanapszám '
                                  f'("{cell_text(r, c_munkanapok)}") — 1 napként számolva'))
            munkanapok = 1
        entry = {
            'nev': cell_text(r, fmt.col('nev')),
            'adoazonosito': cell_text(r, fmt.col('adoazonosito')),
            'taj': cell_text(r, fmt.col('taj')),
            'start_serial': start_serial,
            'munkanapok': munkanapok,
        }
        cutoff = cutoffs.get(entry_withdrawal_key(entry))
        if cutoff and _row_sort_key(r, fmt, 0, i) < cutoff[0]:
            continue  # később visszavonták
        if current_month_serial and serial_to_month_serial(start_serial) > current_month_serial:
            future_entries.append(entry)
        else:
            current_entries.append(entry)

    return current_month_serial, current_entries, future_entries


def merge_entries(current_entries, carried_entries):
    merged = list(current_entries)
    seen = {entry_key(e) for e in current_entries}
    for e in carried_entries:
        if entry_key(e) not in seen:
            seen.add(entry_key(e))
            merged.append(e)
    return merged


# --- snapshot szerializáció ---
#
# A kész munkafüzetet nem tároljuk: a 'generate_output' determinisztikus, ezért
# elég a bemeneteit (header, data_rows, entries, fmt) elmenteni, és letöltéskor
# újragenerálni. Az alábbi függvények ezt a JSON-ra alakítást végzik.

def format_by_key(key):
    """Az InputFormat visszakeresése a kulcsa alapján (ismeretlennél az alap)."""
    for fmt in FORMATS:
        if fmt.key == key:
            return fmt
    return DEFAULT_FORMAT


def _cell_to_json(value):
    """Egy cellaérték JSON-ra alakítása.

    Az openpyxl dátumcellát datetime-ként adja vissza, amit a JSON nem ismer,
    ezért jelölt objektummá alakítjuk. Minden más nem-primitív típusból string
    lesz — a kimenetben úgyis szövegként jelenne meg.
    """
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, datetime):
        return {'__dt__': value.isoformat()}
    return str(value)


def _cell_from_json(value):
    if isinstance(value, dict) and '__dt__' in value:
        return datetime.fromisoformat(value['__dt__'])
    return value


def rows_to_json(rows):
    return [[_cell_to_json(c) for c in row] for row in rows]


def rows_from_json(rows):
    return [tuple(_cell_from_json(c) for c in row) for row in rows]


def entries_to_json(entries):
    """A bejegyzések ISO dátummal — a serial az EPOCH-tól függ, az ISO nem."""
    return [{
        'nev': e['nev'],
        'adoazonosito': e['adoazonosito'],
        'taj': e['taj'],
        'start_date': serial_to_iso(e['start_serial']),
        'munkanapok': e['munkanapok'],
    } for e in entries]


def entries_from_json(items):
    return [{
        'nev': i['nev'],
        'adoazonosito': i['adoazonosito'],
        'taj': i['taj'],
        'start_serial': iso_to_serial(i['start_date']),
        'munkanapok': i['munkanapok'],
    } for i in items]


def build_snapshot(header, data_rows, entries, fmt, prior=None, year=None):
    """A 'generate_output' bemeneteiből JSON-ra alakítható snapshot.

    Az éves táblához használt előző hónapok ('prior') is bekerülnek, így a
    munkafüzet később pontosan ugyanígy állítható elő — akkor is, ha egy
    korábbi hónapot azóta újra feldolgoztak.
    """
    snapshot = {
        'v': 1,
        'fmt': fmt.key,
        'header': [_cell_to_json(c) for c in (header or [])],
        'data_rows': rows_to_json(data_rows),
        'entries': entries_to_json(entries),
    }
    if prior is not None:
        snapshot['prior'] = prior
        snapshot['year'] = year
    return snapshot


def snapshot_to_args(snapshot):
    """A snapshotból a 'generate_output' hívásához szükséges értékek."""
    header = tuple(_cell_from_json(c) for c in snapshot.get('header') or [])
    return {
        'header': header,
        'data_rows': rows_from_json(snapshot.get('data_rows') or []),
        'entries': entries_from_json(snapshot.get('entries') or []),
        'fmt': format_by_key(snapshot.get('fmt')),
        # régebbi snapshotban nincs: ilyenkor a hívó tölti az előző hónapokat
        'prior': snapshot.get('prior'),
        'year': snapshot.get('year'),
    }


def content_hash(path):
    """A munkafüzet tartalmának ujjlenyomata (a metaadatok nélkül).

    A nyers fájl-hash erre nem alkalmas: az openpyxl a létrehozás idejét is
    beleírja a 'docProps/core.xml'-be, így két egyébként azonos generálás
    fájlszinten mindig eltér. Itt csak a munkalapok cellaértékeit hasheljük.
    """
    wb = load_workbook(path)
    h = hashlib.sha256()
    for name in wb.sheetnames:
        h.update(name.encode('utf-8'))
        for row in wb[name].values:
            h.update(repr(row).encode('utf-8'))
    wb.close()
    return h.hexdigest()


MONTH_NAMES = ('Január', 'Február', 'Március', 'Április', 'Május', 'Június',
               'Július', 'Augusztus', 'Szeptember', 'Október', 'November', 'December')


def _person_row_key(adoazonosito, taj, nev):
    ado = ado_key(adoazonosito)
    if ado:
        return 'ado:' + ado
    t = taj_key(taj)
    return 'taj:' + t if t else 'nev:' + (nev or '').strip().lower()


def _write_year_sheet(ws, entries, prior, year):
    """A 'ki hány napot dolgozott' éves tábla a minta szerint.

    1. sor üres; 2. sor: 'név', majd havonta két oszlop (a hónap neve az
    összeg-oszlop fölött), végül 'bérkifizetés' és 'munkanapok'; 3. sor:
    'Nap' / 'Összeg' és 'összesen'. Személyenként két sor: a névsorban a
    havi napok és a munkanapok összege (képlet), alatta az összeg-sor — ezt
    a felhasználó tölti, az összesítő képlet összeadja.
    A hónapokat a nap naptári hónapja szerint számoljuk.
    """
    # (adóazonosító, TAJ, név, napok, havi napszámok) — a korábbi hónapok előbb,
    # a mostani hónap utolsóként (a név innen nyer)
    records = [(p.get('adoazonosito', ''), p.get('taj', ''), p.get('nev', ''),
                [iso_to_serial(d) for d in p.get('dates') or ()], p.get('months') or {})
               for p in prior or ()]
    records += [(e['adoazonosito'], e['taj'], e['nev'],
                 [e['start_serial'] + i for i in range(e['munkanapok'])], {})
                for e in entries]

    # Ugyanaz az ember egy sorba: a csak TAJ-jal vagy csak névvel szereplő
    # bejegyzés az adóazonosítós személyhez kerül, ha a TAJ / a név egyezik.
    by_taj, by_name = {}, {}
    for ado, taj, nev, _, _ in records:
        if ado_key(ado):
            if taj_key(taj).isdigit():
                by_taj.setdefault(taj_key(taj), _person_row_key(ado, '', ''))
            by_name.setdefault(' '.join((nev or '').lower().split()), _person_row_key(ado, '', ''))

    people = {}
    for ado, taj, nev, serials, months in records:
        if ado_key(ado):
            key = _person_row_key(ado, '', '')
        elif taj_key(taj).isdigit() and taj_key(taj) in by_taj:
            key = by_taj[taj_key(taj)]
        else:
            key = by_name.get(' '.join((nev or '').lower().split())) or _person_row_key(ado, taj, nev)
        person = people.setdefault(key, {'nev': nev, 'dates': set(), 'counts': {}})
        person['nev'] = nev or person['nev']
        person['dates'].update(serials)
        for ym_, n in months.items():
            person['counts'][ym_] = n

    if year is None:
        years = [serial_to_date(s).year for p in people.values() for s in p['dates']]
        years += [int(k[:4]) for p in people.values() for k in p['counts']]
        year = max(set(years), key=years.count) if years else datetime.now().year

    first_day_col, last_amount_col = 2, 1 + 2 * 12          # B .. Y
    pay_col, days_col = last_amount_col + 1, last_amount_col + 2   # Z, AA
    ws.append([])
    # a hónap neve az összeg-oszlop fölött van (C, E, …), mint a mintában
    ws.append(['név'] + [x for name in MONTH_NAMES for x in ('', name)]
              + ['bérkifizetés', 'munkanapok'])
    ws.append(['', 'Nap', 'Összeg'] + [''] * (last_amount_col - 3) + ['összesen', 'összesen'])

    bold = Font(bold=True)
    for cell in list(ws[2]) + list(ws[3]):
        cell.font = bold

    for person in sorted(people.values(), key=lambda p: p['nev'].lower()):
        by_month = {}
        for s in person['dates']:
            d = serial_to_date(s)
            if d.year == year:
                by_month[d.month] = by_month.get(d.month, 0) + 1
        for ym_, n in person['counts'].items():      # kézi éves táblából
            if int(ym_[:4]) == year and int(ym_[5:7]) not in by_month:
                by_month[int(ym_[5:7])] = n
        if not by_month:
            continue
        r = ws.max_row + 1
        ws.cell(row=r, column=1).value = person['nev']
        for month, days in by_month.items():
            ws.cell(row=r, column=first_day_col + 2 * (month - 1)).value = days
        last = get_column_letter(last_amount_col)
        ws.cell(row=r, column=days_col).value = f'=SUM(B{r}:{last}{r})'
        ws.cell(row=r + 1, column=pay_col).value = f'=SUM(C{r + 1}:{last}{r + 1})'

    # oszlopszélességek a minta szerint
    longest = max([len('név')] + [len(p['nev']) for p in people.values()])
    ws.column_dimensions['A'].width = max(15, longest + 2)
    for m in range(12):
        ws.column_dimensions[get_column_letter(first_day_col + 2 * m)].width = 4.3
        ws.column_dimensions[get_column_letter(first_day_col + 2 * m + 1)].width = 11
    ws.column_dimensions[get_column_letter(pay_col)].width = 13
    ws.column_dimensions[get_column_letter(days_col)].width = 12


def generate_output(input_path, header, data_rows, entries, output_path=None,
                    fmt=DEFAULT_FORMAT, persons=None, prior=None, year=None):
    """A statisztika-munkafüzet.

    'prior': az év korábbi hónapjainak napjai a 'ki hány napot dolgozott'
    éves táblához — [{'nev', 'adoazonosito', 'taj', 'dates': ['ÉÉÉÉ-HH-NN']}]
    (a mentett előző hónapokból; lásd worklog.prior_year_days). 'year': a
    tábla éve (alapból a feldolgozott napok leggyakoribb éve).
    """
    persons = persons or {}
    by_date = {}
    by_name = {}

    for entry in entries:
        nev = entry['nev']
        if nev not in by_name:
            by_name[nev] = {'adoazonosito': entry['adoazonosito'],
                            'taj': entry['taj'], 'dates': set()}
        for i in range(entry['munkanapok']):
            serial = entry['start_serial'] + i
            by_date.setdefault(serial, []).append({'nev': nev})
            by_name[nev]['dates'].add(serial)

    sorted_dates = sorted(by_date.keys())
    sorted_names = sorted(by_name.keys(), key=lambda x: x.lower())

    wb_out = Workbook()
    wb_out.remove(wb_out.active)

    RED_FILL = PatternFill("solid", fgColor="FF0000")
    BLUE_FILL = PatternFill("solid", fgColor="4472C4")
    WHITE_FONT = Font(color="FFFFFF", bold=True)

    # forrásadat lap (aktuális fájl) – a bemeneti munkalap nevével
    ws_ebev = wb_out.create_sheet(fmt.output_sheet)
    ws_ebev.append(list(header))
    for r in data_rows:
        ws_ebev.append(list(r))
        row_idx = ws_ebev.max_row
        if is_hibas(r, fmt):
            fill, font = RED_FILL, WHITE_FONT
        elif is_torles(r, fmt):
            fill, font = BLUE_FILL, WHITE_FONT
        else:
            continue
        for col in range(1, len(r) + 1):
            cell = ws_ebev.cell(row=row_idx, column=col)
            cell.fill = fill
            cell.font = font

    # Dátum Szerint
    ws_datum = wb_out.create_sheet('Dátum Szerint')
    for serial in sorted_dates:
        workers = by_date[serial]
        ws_datum.append([serial_to_date(serial), workers[0]['nev'], '', '', '', ''])
        for w in workers[1:]:
            ws_datum.append(['', w['nev'], '', '', '', ''])
        ws_datum.append([''] * 6)

    # Név Szerint
    ws_nev = wb_out.create_sheet('Név Szerint')
    for nev in sorted_names:
        info = by_name[nev]
        p = find_person(persons, info['adoazonosito'], info['taj'])[1] or {}
        ws_nev.append(['név:', nev, '', '', ''])
        ws_nev.append(['szül.név', p.get('szul_nev', ''), '', '', ''])
        ws_nev.append(['anyja neve:', p.get('anya_neve', ''), '', '', ''])
        ws_nev.append(['szül.hely, idő:', p.get('szul_hely_ido', ''), '', '', ''])
        ws_nev.append(['adóazonosító:', info['adoazonosito'], '', '', ''])
        ws_nev.append(['TAJ-szám:', normalize_taj(info['taj']), '', '', ''])
        ws_nev.append(['lakcím:', p.get('lakcim', ''), '', '', ''])
        ws_nev.append([''] * 5)
        unique_dates = sorted(info['dates'])
        for serial in unique_dates:
            ws_nev.append([serial_to_date(serial), 1, '', '', ''])
        ws_nev.append(['', len(unique_dates), '', '', ''])
        ws_nev.append([''] * 5)

    # Ki hány napot dolgozott — éves tábla (január–december), az előző
    # hónapok napjaival együtt (lásd _write_year_sheet)
    ws_hany = wb_out.create_sheet('ki hány napot dolgozott')
    _write_year_sheet(ws_hany, entries, prior, year)

    for ws in wb_out.worksheets:
        if ws is not ws_hany:
            autofit(ws)

    if output_path is None:
        import os
        output_path = os.path.splitext(input_path)[0] + '_statisztika.xlsx'
    wb_out.save(output_path)
    return output_path
