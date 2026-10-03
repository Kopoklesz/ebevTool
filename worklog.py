"""Munkanapló: személyenként mely napokon dolgozott, melyik cégnél.

Forrás a havi statisztikák felhőben mentett tartalma (snapshot), cégenként és
hónaponként CSAK A LEGUTOLSÓ feldolgozásé — így ha egy javított fájlt töltöttek
fel újra (pl. kihúztak belőle valakit), a korábbi verzió nem számít. A napok
pontosan azok, amik a havi statisztikába is bekerültek (generate_output: a
kezdő naptól 'munkanapok' egymást követő nap).

Ahol egy hónapnak nincs snapshotja (a 2026.10.01 előtti feldolgozások), ott
a helyi archívumból visszamenőleg betöltött adat pótolhatja (archive_import).

A visszavonásokat cégszinten, a hónapokon átívelve is alkalmazzuk: ha egy
későbbi havi fájl vonja vissza egy korábbi hónap bejelentését, a napló már nem
számolja (a régi havi statisztika ettől nem változik). Az archívumból
betöltött hónapoknál nincs bejelentés-szintű adat, ott ez nem alkalmazható.
"""
from datetime import date

import generate
from normalize import normalize_taj


def latest_processing(history_rows):
    """{hónap: előzmény-sor} — hónaponként a legutolsó feldolgozás.

    Akkor is ez számít, ha nincs mentett tartalma (pl. a mentés nem sikerült,
    vagy egy régi verziójú gép dolgozta fel): ilyenkor NEM esünk vissza egy
    korábbi, elavult verzióra — a hónap a hívónál hiányzóként jelenik meg
    (és az archívumból betöltött adat pótolhatja)."""
    latest = {}
    for row in history_rows:
        ym = row.get('year_month')
        if not ym:
            continue
        if ym not in latest or (row.get('processed_at') or '') > (latest[ym].get('processed_at') or ''):
            latest[ym] = row
    return latest


def person_key(adoazonosito, taj, nev):
    """A személy azonosítója a naplóban: adóazonosító, ennek híján TAJ, végül név."""
    ado = generate.ado_key(adoazonosito)
    if ado:
        return 'ado:' + ado
    t = generate.taj_key(taj)
    if t:
        return 'taj:' + t
    return 'nev:' + (nev or '').strip().lower()


def _serial_to_day(serial):
    return generate.serial_to_date(serial).date()


def build(companies):
    """A munkanapló felépítése.

    'companies': {cég: {'snapshots': {hónap: snapshot-payload},
                        'imports': {hónap: {'source', 'persons': [...]}}}}

    Visszatérés: {személykulcs: {
        'nev', 'adoazonosito', 'taj',
        'companies': {cég: {'days': {date: hónap (melyik havi statisztikában)},
                            'archive_months': set(hónap)}},
        'withdrawn': [(cég, kezdő nap (date), név)],
    }}
    """
    people = {}

    def person(adoazonosito, taj, nev):
        key = person_key(adoazonosito, taj, nev)
        p = people.setdefault(key, {'nev': nev, 'adoazonosito': adoazonosito, 'taj': taj,
                                    'companies': {}, 'withdrawn': []})
        # a legfrissebb (később feldolgozott) adat nyer a megjelenített mezőknél
        p['nev'] = nev or p['nev']
        p['adoazonosito'] = adoazonosito or p['adoazonosito']
        p['taj'] = taj or p['taj']
        return p

    for company, data in companies.items():
        snapshots = data.get('snapshots') or {}
        imports = data.get('imports') or {}

        # cégszintű visszavonások, az összes hónap sorai alapján
        events = []
        parsed = {}
        for order, ym in enumerate(sorted(snapshots)):
            args = generate.snapshot_to_args(snapshots[ym])
            parsed[ym] = args
            events.extend(generate.row_events(args['data_rows'], args['fmt'], order=order))
        valid = generate.valid_registrations(events)

        for ym in sorted(snapshots):
            for entry in parsed[ym]['entries']:
                p = person(entry['adoazonosito'], entry['taj'], entry['nev'])
                key = generate.entry_withdrawal_key(entry)
                if key in valid and entry['munkanapok'] not in valid[key]:
                    item = (company, _serial_to_day(entry['start_serial']), entry['nev'])
                    if item not in p['withdrawn']:
                        p['withdrawn'].append(item)
                    continue
                c = p['companies'].setdefault(company, {'days': {}, 'archive_months': set()})
                for i in range(entry['munkanapok']):
                    c['days'].setdefault(_serial_to_day(entry['start_serial'] + i), ym)

        for ym in sorted(imports):
            if ym in snapshots:
                continue  # a valódi, frissebb tartalom erősebb
            for rec in imports[ym].get('persons') or []:
                p = person(rec.get('adoazonosito', ''), rec.get('taj', ''), rec.get('nev', ''))
                c = p['companies'].setdefault(company, {'days': {}, 'archive_months': set()})
                c['archive_months'].add(ym)
                for iso in rec.get('dates') or []:
                    c['days'].setdefault(date.fromisoformat(iso), ym)

    # akinek minden napját visszavonták, az is maradjon látható (üres céggel)
    return people


def find(people, adoazonosito, taj, nev=''):
    """A személy naplóbejegyzései: adóazonosító szerint, tartalékként TAJ.

    Több bejegyzést is visszaadhat (pl. egyik helyen adóval, másikon csak
    TAJ-jal rögzítve); ezeket a hívó összevonja (merge)."""
    ado, t = generate.ado_key(adoazonosito), generate.taj_key(taj)
    found = []
    for key, p in people.items():
        p_ado = generate.ado_key(p['adoazonosito'])
        if ado and p_ado == ado:
            found.append(p)
        elif t and generate.taj_key(p['taj']) == t and (not p_ado or not ado):
            found.append(p)
        elif not ado and not t and nev and key == person_key('', '', nev):
            found.append(p)
    return found


def merge(records):
    """Több naplóbejegyzés összevonása egybe (azonos nap egyszer számít)."""
    out = {'nev': '', 'adoazonosito': '', 'taj': '', 'companies': {}, 'withdrawn': []}
    for r in records:
        out['nev'] = out['nev'] or r['nev']
        out['adoazonosito'] = out['adoazonosito'] or r['adoazonosito']
        out['taj'] = out['taj'] or r['taj']
        for company, c in r['companies'].items():
            oc = out['companies'].setdefault(company, {'days': {}, 'archive_months': set()})
            for d, ym in c['days'].items():
                oc['days'].setdefault(d, ym)
            oc['archive_months'] |= c['archive_months']
        out['withdrawn'] += [w for w in r['withdrawn'] if w not in out['withdrawn']]
    return out


def summarize(record):
    """Megjelenítéshez: cégenként napok száma, első/utolsó nap, havi bontás.

    Visszatérés: {'companies': [{'company', 'count', 'first', 'last',
    'months': [(hónap 'ÉÉÉÉ-HH', [date, ...])], 'archive_months'}],
    'total_days', 'total_unique_days', 'first', 'last'}
    A 'total_days' a cégenkénti napok összege; ha valaki ugyanazon a napon két
    cégnél is dolgozott, a 'total_unique_days' ezt egyszer számolja.
    """
    companies = []
    all_days = set()
    total = 0
    for company in sorted(record['companies'], key=str.lower):
        c = record['companies'][company]
        days = sorted(c['days'])
        if not days:
            continue
        months = {}
        for d in days:
            months.setdefault(f'{d.year:04d}-{d.month:02d}', []).append(d)
        companies.append({'company': company, 'count': len(days), 'first': days[0],
                          'last': days[-1], 'months': sorted(months.items()),
                          'archive_months': sorted(c['archive_months'])})
        all_days.update(days)
        total += len(days)
    return {'companies': companies, 'total_days': total, 'total_unique_days': len(all_days),
            'first': min(all_days) if all_days else None,
            'last': max(all_days) if all_days else None}


def export_xlsx(record, summary, path):
    """A munkanapló Excelbe: összesítő lap + napok listája lap."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    bold = Font(bold=True)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Összesítés'
    for label, value in (('Név', record['nev']), ('Adóazonosító', record['adoazonosito']),
                         ('TAJ-szám', normalize_taj(record['taj'])),
                         ('Összes munkanap', summary['total_days'])):
        ws.append([label, value])
        ws.cell(row=ws.max_row, column=1).font = bold
    ws.append([])
    ws.append(['Cég', 'Napok', 'Első nap', 'Utolsó nap', 'Megjegyzés'])
    for cell in ws[ws.max_row]:
        cell.font = bold
    for c in summary['companies']:
        note = ('archívumból: ' + ', '.join(c['archive_months'])) if c['archive_months'] else ''
        ws.append([c['company'], c['count'], c['first'], c['last'], note])
    if record['withdrawn']:
        ws.append([])
        ws.append(['Visszavont bejelentések (nem számítanak bele)'])
        ws.cell(row=ws.max_row, column=1).font = bold
        for company, start, _ in sorted(record['withdrawn']):
            ws.append([company, start])

    days = wb.create_sheet('Napok')
    days.append(['Cég', 'Dátum', 'Hónap'])
    for cell in days[1]:
        cell.font = bold
    for c in summary['companies']:
        for ym, ds in c['months']:
            for d in ds:
                days.append([c['company'], d, ym])
    for sheet in (ws, days):
        generate.autofit(sheet)
    wb.save(path)


def counts_for_year(stat_ym, ym):
    """Kell-e a 'stat_ym' havi statisztika az 'ym' évének táblájához: az év
    'ym' előtti hónapjai, és az előző év decembere (abból átnyúlhat januárra)."""
    year = int(ym[:4])
    return stat_ym < ym and (stat_ym[:4] == ym[:4] or stat_ym == f'{year - 1}-12')


def prior_year_days(company, snapshots, imports, ym, year_tables=None, resolve=None):
    """Az év korábbi hónapjainak napjai a 'ki hány napot dolgozott' táblához.

    'snapshots' / 'imports': a cég hónaponkénti mentett tartalmai, illetve az
    archívumból betöltött hónapjai (bármelyik hónapra; itt szűrünk). Csak az
    'ym' évének a 'ym'-nél KORÁBBI hónapjai számítanak — a mostani hónapot a
    feldolgozás maga adja.
    Azokra a hónapokra, amelyekhez nincs havi statisztika, a kézi
    munkafüzetekből betöltött éves tábla ('year_tables': {év: {'persons': [...
    'months': {hónap: napok}]}}) napszámai kerülnek; ahol ez sincs, üres marad.
    'resolve(név) -> (adóazonosító, TAJ)': az éves táblában csak névvel
    szereplők azonosítása (hogy ugyanabba a sorba kerüljenek).
    Visszatérés: [{'nev', 'adoazonosito', 'taj', 'dates': ['ÉÉÉÉ-HH-NN']}
    vagy {..., 'months': {'ÉÉÉÉ-HH': napok}}].
    """
    year = ym[:4]
    wanted = lambda m: counts_for_year(m, ym)
    data = {company: {
        'snapshots': {m: v for m, v in (snapshots or {}).items() if wanted(m)},
        'imports': {m: v for m, v in (imports or {}).items() if wanted(m)},
    }}
    out = []
    for person in build(data).values():
        days = (person['companies'].get(company) or {}).get('days') or {}
        dates = sorted(d.isoformat() for d, stat_ym in days.items()
                       if wanted(stat_ym) and str(d.year) == year)
        if dates:
            out.append({'nev': person['nev'], 'adoazonosito': person['adoazonosito'],
                        'taj': person['taj'], 'dates': dates})

    covered = set(data[company]['snapshots']) | set(data[company]['imports'])
    table = (year_tables or {}).get(year) or (year_tables or {}).get(int(year)) or {}
    for p in table.get('persons') or ():
        months = {}
        for m, n in (p.get('months') or {}).items():
            m_ym = f'{year}-{int(m):02d}'
            if m_ym < ym and m_ym not in covered and n:
                months[m_ym] = int(n)
        if months:
            ado, taj = p.get('adoazonosito', ''), p.get('taj', '')
            if not ado and resolve:
                ado, taj = resolve(p.get('nev', '')) or (ado, taj)
            out.append({'nev': p.get('nev', ''), 'adoazonosito': ado, 'taj': taj,
                        'months': months})
    return out
